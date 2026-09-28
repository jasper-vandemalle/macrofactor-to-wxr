"""Reads workouts from a MacroFactor XLSX export."""

import re
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from mf2wxr.model import Exercise, Set, Workout

WORKOUT_SHEET = "Workout Log"
BODYWEIGHT_SHEET = "Quick Export"

# header in the Workout Log sheet -> our key
WORKOUT_COLUMNS = {
    "date": "date",
    "workout": "workout_name",
    "exercise": "exercise",
    "weight (kg)": "weight",
    "reps": "reps",
    "set type": "set_type",
    "distance short (m)": "distance_short",
    "distance long (km)": "distance_long",
    "rir": "rir",
}
REQUIRED_COLUMNS = ("date", "exercise", "weight", "reps")


class ExportError(ValueError):
    """The file isn't a MacroFactor export we can read."""


def read_workout(path: str | Path, on: date | None = None) -> Workout:
    """Read the workout logged on the given date, or the most recent one.

    The bodyweight is the most recent weigh-in on or before the workout date.
    """
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        workout = _read_workout(wb, on)
        workout.bodyweight_kg = _read_bodyweight(wb, workout.date)
        return workout
    finally:
        wb.close()


def available_dates(path: str | Path) -> list[date]:
    """All dates with a logged workout, oldest first."""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        columns, rows = _workout_rows(wb)
        return sorted({d for row in rows if (d := _to_date(row[columns["date"]])) is not None})
    finally:
        wb.close()


def _workout_rows(wb) -> tuple[dict[str, int], list[tuple]]:
    if WORKOUT_SHEET not in wb.sheetnames:
        raise ExportError(f"No '{WORKOUT_SHEET}' sheet found; is this a MacroFactor workout export?")

    sheet = wb[WORKOUT_SHEET]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    columns = {}
    for i, name in enumerate(header):
        key = WORKOUT_COLUMNS.get(str(name).lower()) if name else None
        if key:
            columns[key] = i

    missing = [c for c in REQUIRED_COLUMNS if c not in columns]
    if missing:
        raise ExportError(f"Missing column(s) in '{WORKOUT_SHEET}': {', '.join(missing)}")

    return columns, list(sheet.iter_rows(min_row=2, values_only=True))


def _read_workout(wb, on: date | None) -> Workout:
    columns, rows = _workout_rows(wb)

    def cell(row, key):
        return row[columns[key]] if key in columns else None

    dates = {d for row in rows if (d := _to_date(cell(row, "date"))) is not None}
    if not dates:
        raise ExportError("No workouts found in the export")

    target = on or max(dates)
    if target not in dates:
        raise ExportError(f"No workout logged on {target.isoformat()}")

    workout_name = None
    # exercises keyed by name, in order of first appearance; supersets interleave their rows
    exercises: dict[str, Exercise] = {}

    for row in rows:
        if _to_date(cell(row, "date")) != target:
            continue

        if workout_name is None:
            workout_name = cell(row, "workout_name")

        set_type = cell(row, "set_type")
        if set_type and "warm" in str(set_type).lower():
            continue

        name = cell(row, "exercise")
        if not name:
            continue
        name = _strip_superset_marker(name)
        exercise = exercises.setdefault(name, Exercise(name))

        weight = cell(row, "weight")
        reps = cell(row, "reps")
        rir = cell(row, "rir")
        exercise.sets.append(
            Set(
                weight_kg=float(weight) if weight else 0.0,
                reps=int(reps) if reps else 0,
                distance_m=_distance_m(cell(row, "distance_short"), cell(row, "distance_long")),
                rpe=10 - float(rir) if rir is not None else None,
            )
        )

    return Workout(
        date=target,
        name=workout_name or "Workout",
        exercises=[_adjust_bilateral(e) for e in exercises.values()],
    )


def _read_bodyweight(wb, on_or_before: date) -> float | None:
    if BODYWEIGHT_SHEET not in wb.sheetnames:
        return None

    sheet = wb[BODYWEIGHT_SHEET]
    header = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), None)
    if not header:
        return None

    date_col = weight_col = None
    for i, name in enumerate(header):
        name = str(name).lower() if name else ""
        if "date" in name:
            date_col = i
        if "weight" in name and "kg" in name:
            weight_col = i
    if date_col is None or weight_col is None:
        return None

    latest = None
    for row in sheet.iter_rows(min_row=2, values_only=True):
        day, weight = _to_date(row[date_col]), row[weight_col]
        if day and weight and day <= on_or_before and (latest is None or day >= latest[0]):
            latest = (day, float(weight))

    return latest[1] if latest else None


def _strip_superset_marker(name: str) -> str:
    """'Bench Press (Barbell) ∈ SS1' -> 'Bench Press (Barbell)'"""
    return re.sub(r"\s*∈\s*SS\d+$", "", name)


def _distance_m(short_m, long_km) -> float | None:
    # weight-over-distance exercises (farmer's carry, sled push) record a distance instead of reps
    if short_m:
        return float(short_m)
    if long_km:
        return float(long_km) * 1000
    return None


def _adjust_bilateral(exercise: Exercise) -> Exercise:
    """MacroFactor records the combined load of two dumbbells (and both sides of a cable fly).

    Single-limb and offset exercises are already recorded per side.
    """
    name = exercise.name.lower()
    single_limb = "single arm" in name or "single leg" in name
    offset = "offset" in name
    if ("dumbbell" in name or "fly" in name) and not single_limb and not offset:
        for s in exercise.sets:
            s.weight_kg /= 2
    return exercise


def _to_date(value) -> date | None:
    if value is None or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    try:
        return date.fromisoformat(str(value)[:10])
    except ValueError:
        return None
