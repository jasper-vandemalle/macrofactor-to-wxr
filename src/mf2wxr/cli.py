import argparse
import sys
from datetime import date
from pathlib import Path

import httpx
from openpyxl import load_workbook

from mf2wxr import auth
from mf2wxr.api import ApiError, WxrApi
from mf2wxr.config import Config, config_dir
from mf2wxr.exercises import AliasError, load_aliases, match_exercises
from mf2wxr.macrofactor import WORKOUT_SHEET, ExportError, read_workout
from mf2wxr.rows import MissingBodyweight, to_rows
from mf2wxr.text import to_text


class UsageError(RuntimeError):
    pass


def main(argv: list[str] | None = None) -> int:
    args = _parser().parse_args(argv)
    config = Config.load()
    store = auth.TokenStore(config_dir() / "token.json")

    with httpx.Client(timeout=30) as http:
        try:
            if args.command == "login":
                return login(args, config, store, http)
            if args.command == "logout":
                store.clear()
                print("Signed out.")
                return 0
            if args.command == "text":
                return text(args, config)
            api = WxrApi(http, lambda: auth.access_token(store, http))
            return push(args, config, api)
        except (UsageError, ExportError, AliasError, MissingBodyweight, auth.AuthError, ApiError, httpx.HTTPError) as e:
            print(f"error: {e}", file=sys.stderr)
            return 1


def login(args, config: Config, store: auth.TokenStore, http: httpx.Client) -> int:
    client_id = args.client_id or config.client_id
    if not client_id:
        raise UsageError(
            "No client id. Register a service under weightxreps settings → Developer API, then pass its id "
            "with --client-id, set MF2WXR_CLIENT_ID, or add `client_id` to the config file"
        )
    store.save(auth.login(client_id, http))
    print("Signed in to weightxreps.")
    return 0


def text(args, config: Config) -> int:
    workout = read_workout(find_export(args.path, config), args.date)
    output = to_text(workout)
    print(output)
    if args.copy:
        _copy(output)
    return 0


def push(args, config: Config, api: WxrApi, aliases_path: Path | None = None) -> int:
    export = find_export(args.path, config)
    workout = read_workout(export, args.date)
    day = api.journal_day(workout.date)

    aliases = load_aliases(aliases_path or config_dir() / "aliases.toml")
    matches = match_exercises([e.name for e in workout.exercises], day.exercises, aliases)
    rows = to_rows(workout, matches)
    new = sorted({m.name for m in matches.values() if m.eid is None})

    print(f"From {export.name}:\n")
    print(to_text(workout))
    if new:
        print("New on weightxreps: " + ", ".join(new))

    clash = day.has_log and not args.replace
    if args.dry_run:
        if clash:
            print(f"{workout.date.isoformat()} already has a log on weightxreps; pushing it needs --replace.")
        print("Dry run, nothing saved.")
        return 0
    if clash:
        raise UsageError(
            f"{workout.date.isoformat()} already has a log on weightxreps. Saving replaces the whole day; "
            "pass --replace if that's what you want"
        )

    api.save(rows, workout.date)
    total_sets = sum(len(e.sets) for e in workout.exercises)
    print(
        f"Logged {workout.name} on {workout.date.isoformat()}: {len(workout.exercises)} exercises, {total_sets} sets."
    )
    return 0


def find_export(path: Path | None, config: Config) -> Path:
    """The given file, or the newest MacroFactor export with workouts in the given (or configured) directory."""
    path = path or config.exports or Path.cwd()
    if path.is_file():
        return path
    if not path.is_dir():
        raise UsageError(f"{path} doesn't exist")

    candidates = sorted(path.glob("MacroFactor*.xlsx"), key=lambda p: p.stat().st_mtime, reverse=True)
    for candidate in candidates:
        if _has_workouts(candidate):
            return candidate
    raise UsageError(f"No MacroFactor workout export in {path}")


def _has_workouts(path: Path) -> bool:
    wb = load_workbook(path, read_only=True)
    try:
        return WORKOUT_SHEET in wb.sheetnames
    finally:
        wb.close()


def _copy(output: str) -> None:
    try:
        import pyperclip
    except ImportError:
        raise UsageError("--copy needs the clipboard extra: install mf2wxr[clipboard]") from None
    pyperclip.copy(output)
    print("Copied to the clipboard.")


def _parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(prog="mf2wxr", description="Push MacroFactor workouts to weightxreps.net")
    commands = parser.add_subparsers(dest="command", required=True)

    login = commands.add_parser("login", help="sign in to weightxreps in the browser")
    login.add_argument("--client-id", help="id of the service you registered on weightxreps")

    commands.add_parser("logout", help="forget the stored weightxreps tokens")

    def add_export_args(command: argparse.ArgumentParser):
        command.add_argument(
            "path",
            nargs="?",
            type=Path,
            help="a MacroFactor export, or a directory to take the newest one from "
            "(default: `exports` from the config file, else the current directory)",
        )
        command.add_argument("--date", type=date.fromisoformat, help="workout date, YYYY-MM-DD (default: the latest)")

    push = commands.add_parser("push", help="log a workout in your weightxreps journal")
    add_export_args(push)
    push.add_argument("--dry-run", action="store_true", help="show what would be logged without saving")
    push.add_argument("--replace", action="store_true", help="overwrite the day if it already has a log")

    text = commands.add_parser("text", help="print the workout as weightxreps editor text")
    add_export_args(text)
    text.add_argument("--copy", action="store_true", help="also copy it to the clipboard")

    return parser


if __name__ == "__main__":
    sys.exit(main())
